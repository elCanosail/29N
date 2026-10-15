#!/usr/bin/env python3
"""Normaliza los resultados oficiales del Interior a esquema comun + VALIDACION D'Hondt (29N).

Entrada  (data/raw/interior/):
    Elecciones-Congreso.xlsx                (escrutinio provincial por candidatura)
    Elecciones-Congreso-cargos-electos.xlsx (cargos electos, definitivo)

Salida (data/processed/interior/):
    results_provincial.jsonl   una linea por (eleccion, provincia, candidatura)
    elections.json             metadatos + agregados por eleccion y partido
    partidos.json              diccionario de partidos/aliases efectivamente usado
    validation.json            resultado de las validaciones obligatorias

VALIDACIONES (por eleccion)
    (a) suma de escaños por partido (provincias) == columna "Total Escaños Nacional"
    (b) D'Hondt sobre votos provinciales reproduce los escaños oficiales publicados
    (c) coincidencia con el fichero de cargos electos (definitivo, independiente)

Nota sobre el voto CERA (residentes ausentes): el fichero publica filas "Votos CERA (..)"
separadas. El reparto de escaños publicado se reproduce con los votos de la
circunscripcion *sin* CERA (verificado contra el fichero de cargos electos); el CERA se
conserva como campo aparte. Ver docs/notes/interior-ingest.md.

No hace commit ni push.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import json
import os
import re
from collections import defaultdict

import openpyxl

from dhondt import dhondt

HERE = os.path.dirname(os.path.abspath(__file__))
CONGRESO_XLSX = "Elecciones-Congreso.xlsx"
CARGOS_XLSX = "Elecciones-Congreso-cargos-electos.xlsx"

# "Votos (NOMBRE - SIGLA)", "Escaños (NOMBRE - SIGLA)", y variantes "... CERA"
ROW_RE = re.compile(r"^(?P<kind>Votos|Escaños)(?P<cera> CERA)?\s*\((?P<lab>.*)\)\s*$")
METRICAS = {
    "Electores", "Electores CER", "Votantes", "Votos válidos", "Votos nulos",
    "Votos en blanco", "Votos a candidaturas", "Escaños",
    "Votantes CERA", "Votos válidos CERA", "Votos nulos CERA",
    "Votos en blanco CERA", "Votos a candidaturas CERA",
}


# ---------------------------------------------------------------- utilidades

def norm_sigla(s):
    return re.sub(r"\s+", " ", (s or "").strip()).upper()


def slug(s):
    s = re.sub(r"[^0-9A-Za-zÁÉÍÓÚÜÑáéíóúüñ+]+", "-", (s or "").strip().lower())
    return re.sub(r"-+", "-", s).strip("-") or "desconocido"


def split_label(lab):
    """'NOMBRE - SIGLA' -> (nombre, sigla). Si no hay ' - ', nombre=sigla=lab."""
    lab = lab.strip()
    if " - " in lab:
        nombre, sigla = lab.rsplit(" - ", 1)
        return nombre.strip(), sigla.strip()
    return lab, lab


def label_candidates(lab):
    """Sufijos del etiqueta empezando en cada ' - ', de mas largo a mas corto.

    El fichero del Interior etiqueta las coaliciones como 'NOMBRE - COALICION'
    donde la propia sigla de la coalicion puede contener ' - ' (p.ej.
    'SUMAR - EN COMU PODEM - SUMAR - ECP'). Probamos todos los sufijos contra el
    diccionario de aliases para resolver la candidatura de forma robusta.
    """
    parts = lab.strip().split(" - ")
    return [" - ".join(parts[i:]) for i in range(len(parts))]


def resolve_party(lab, alias_idx):
    """Id canonico del partido para una etiqueta de candidatura, o None."""
    for cand in label_candidates(lab):
        pid = alias_idx.get(norm_sigla(cand))
        if pid:
            return pid
    return None


def party_key(lab, alias_idx):
    """Clave canonica estable para comparar candidaturas entre ficheros."""
    pid = resolve_party(lab, alias_idx)
    if pid:
        return pid
    return norm_sigla(lab.rsplit(" - ", 1)[-1])


def load_alias_index(path):
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    idx = {}
    for pid, meta in data["partidos"].items():
        for a in [meta.get("nombre", "")] + meta.get("aliases", []):
            if a:
                idx[norm_sigla(a)] = pid
    return idx


def header_index(rows):
    for i, r in enumerate(rows[:10]):
        if r and r[0] == "Fecha":
            return i
    raise SystemExit("No encuentro la fila de cabecera ('Fecha').")


def province_cols(hdr):
    cols = []
    for i in range(4, len(hdr)):
        v = hdr[i]
        if not v:
            continue
        m = re.match(r"^(\d+)\s*-\s*(.+)$", str(v))
        if m:
            cols.append((i, int(m.group(1)), m.group(2).strip()))
    return cols


# ---------------------------------------------------------------- parsing

def parse_congreso(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hi = header_index(rows)
    hdr = rows[hi]
    provs = province_cols(hdr)
    natcol = next((i for i, v in enumerate(hdr)
                   if v and str(v).strip().lower().startswith("total esca")), None)
    by_el = defaultdict(list)
    for r in rows[hi + 1:]:
        if r[0] is None:
            continue
        by_el[str(r[0])[:10]].append(r)
    return provs, natcol, by_el


def parse_cargos(path):
    """(fecha, prov_code) -> {sigla_norm: n_electos}"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hi = header_index(rows)
    hdr = rows[hi]
    col = {str(v).strip(): i for i, v in enumerate(hdr) if v}
    ci_pc, ci_sig = col["Código Provincia"], col["Siglas Candidatura"]
    out = defaultdict(lambda: defaultdict(int))
    for r in rows[hi + 1:]:
        if r[0] is None:
            continue
        out[(str(r[0])[:10], str(r[ci_pc]))][norm_sigla(str(r[ci_sig]))] += 1
    return out


def parse_election_block(rr, provs, natcol):
    """Extrae votos/esc anos/metricas de las filas de una eleccion."""
    plain = defaultdict(dict)   # label -> {prov: votos}
    cera = defaultdict(dict)    # label -> {prov: votos CERA}
    esc = defaultdict(dict)     # label -> {prov: escaños}
    nac = {}                    # label -> escaños nacionales
    metrics = defaultdict(dict)
    summary = None
    for r in rr:
        desc = str(r[3]) if r[3] else ""
        if desc == "Escaños":
            summary = r
            continue
        m = ROW_RE.match(desc)
        if m:
            lab, kind, is_cera = m.group("lab").strip(), m.group("kind"), bool(m.group("cera"))
            tgt = cera if is_cera else plain
            for (ci, code, _name) in provs:
                v = r[ci]
                if not v:
                    continue
                if kind == "Votos":
                    tgt[lab][code] = tgt[lab].get(code, 0) + int(v)
                elif not is_cera:
                    esc[lab][code] = esc[lab].get(code, 0) + int(v)
            if kind == "Escaños" and not is_cera and natcol is not None and r[natcol]:
                # una misma candidatura puede aparecer en varias filas (subconjuntos de
                # provincias); la columna nacional es el subtotal de cada fila -> se acumula
                nac[lab] = nac.get(lab, 0) + int(r[natcol])
        elif desc in METRICAS:
            for (ci, code, _name) in provs:
                if r[ci] is not None:
                    metrics[desc][code] = r[ci]
    prov_seats = {code: (int(summary[ci]) if summary and summary[ci] else 0)
                  for (ci, code, _name) in provs}
    return plain, cera, esc, nac, metrics, prov_seats


# ---------------------------------------------------------------- normalizacion

def normalize_election(date, provs, plain, cera, esc, prov_seats, alias_idx, partidos_used):
    records = []
    prov_votes = defaultdict(dict)
    prov_esc = defaultdict(dict)
    for lab in set(plain) | set(cera) | set(esc):
        nombre, sigla = split_label(lab)
        key = norm_sigla(sigla)
        pid = resolve_party(lab, alias_idx) or slug(sigla)
        partidos_used[pid]["siglas"].add(sigla)
        partidos_used[pid]["nombres"].add(nombre)
        partidos_used[pid]["auto"] = partidos_used[pid]["auto"] and (resolve_party(lab, alias_idx) is None)
        for (ci, code, pname) in provs:
            v = plain.get(lab, {}).get(code, 0)
            vc = cera.get(lab, {}).get(code, 0)
            se = esc.get(lab, {}).get(code, 0)
            if v or vc or se:
                records.append({
                    "eleccion": date, "provincia": code, "provincia_nombre": pname,
                    "candidatura": lab, "sigla": sigla, "nombre": nombre,
                    "partido_id": pid, "votos": v, "votos_cera": vc, "escanos": se,
                })
            if v:
                prov_votes[code][lab] = v
            if se:
                prov_esc[code][lab] = se
    return records, prov_votes, prov_esc


# ---------------------------------------------------------------- validacion

def validate(date, provs, prov_votes, prov_esc, prov_seats, metrics, nac, cargos, alias_idx):
    val = {"id": date, "escanos_totales": sum(prov_seats.values()), "provincias": {},
           "a_sum_escanos_ok": True, "b_dhondt_ok": True, "c_cargos_ok": True, "fallos": []}

    # (a) suma de escanos provinciales == nacional (por partido)
    esc_totals = defaultdict(int)
    for _code, dd in prov_esc.items():
        for lab, s in dd.items():
            esc_totals[lab] += s
    for lab in set(esc_totals) | set(nac):
        s = esc_totals.get(lab, 0)
        n = nac.get(lab, 0)
        if s != n:
            val["a_sum_escanos_ok"] = False
            val["fallos"].append({"tipo": "a", "candidatura": lab,
                                  "suma_provincias": s, "nacional": n})

    # (b) D'Hondt por provincia
    for (ci, code, pname) in provs:
        seats = prov_seats[code]
        votes = dict(prov_votes.get(code, {}))
        off = dict(prov_esc.get(code, {}))
        blank = metrics.get("Votos en blanco", {}).get(code, 0) or 0
        valid = sum(votes.values()) + blank
        calc_nz = {c: s for c, s in dhondt(votes, seats, 0.03, valid).items() if s}
        off_nz = {c: s for c, s in off.items() if s}
        ok = calc_nz == off_nz
        val["provincias"][str(code)] = {"nombre": pname, "escanos": seats, "ok": ok,
                                        "blanco": blank, "validos": valid}
        if not ok:
            val["b_dhondt_ok"] = False
            diff = {k: off_nz.get(k, 0) - calc_nz.get(k, 0) for k in set(off_nz) | set(calc_nz)
                    if off_nz.get(k, 0) != calc_nz.get(k, 0)}
            val["fallos"].append({"tipo": "b", "provincia": code, "provincia_nombre": pname,
                                  "escanos": seats, "oficial": off_nz, "dhondt": calc_nz, "dif": diff})

    # (c) cargos electos (definitivo) == escrutinio por provincia (clave canonica)
    for (ci, code, pname) in provs:
        cl = cargos.get((date, str(code)))
        if cl is None:
            continue
        escrut = defaultdict(int)
        for lab, s in prov_esc.get(code, {}).items():
            if s:
                escrut[party_key(lab, alias_idx)] += s
        car = defaultdict(int)
        for sig, n in cl.items():
            car[party_key(sig, alias_idx)] += n
        if dict(car) != dict(escrut):
            val["c_cargos_ok"] = False
            val["fallos"].append({"tipo": "c", "provincia": code, "provincia_nombre": pname,
                                  "cargos": dict(car), "escrutinio": dict(escrut)})

    val["resumen"] = {"escanos": val["escanos_totales"],
                      "provincias_ok": sum(1 for p in val["provincias"].values() if p["ok"]),
                      "provincias_total": len(val["provincias"])}
    return val


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--raw", default="data/raw/interior")
    ap.add_argument("--out", default="data/processed/interior")
    args = ap.parse_args()

    alias_idx = load_alias_index(os.path.join(HERE, "partidos_aliases.json"))
    provs, natcol, by_el = parse_congreso(os.path.join(args.raw, CONGRESO_XLSX))
    cargos = parse_cargos(os.path.join(args.raw, CARGOS_XLSX))
    os.makedirs(args.out, exist_ok=True)

    partidos_used = defaultdict(lambda: {"siglas": set(), "nombres": set(), "auto": True})
    elections, validations = [], []
    n_rec = 0

    with open(os.path.join(args.out, "results_provincial.jsonl"), "w", encoding="utf-8") as recs_f:
        for date in sorted(by_el):
            plain, cera, esc, nac, metrics, prov_seats = parse_election_block(
                by_el[date], provs, natcol)
            records, prov_votes, prov_esc = normalize_election(
                date, provs, plain, cera, esc, prov_seats, alias_idx, partidos_used)
            for rec in records:
                recs_f.write(json.dumps(rec, ensure_ascii=False) + "\n")
            n_rec += len(records)

            val = validate(date, provs, prov_votes, prov_esc, prov_seats, metrics, nac, cargos,
                           alias_idx)
            validations.append(val)

            pt = defaultdict(lambda: {"votos": 0, "escanos": 0})
            for lab in set(plain) | set(esc):
                nombre, sigla = split_label(lab)
                pid = resolve_party(lab, alias_idx) or slug(sigla)
                pt[pid]["votos"] += sum(plain.get(lab, {}).values())
                pt[pid]["escanos"] += sum(esc.get(lab, {}).values())
            elections.append({
                "id": date, "fecha": date, "tipo": "congreso",
                "camara": "Congreso de los Diputados",
                "escanos_totales": val["escanos_totales"],
                "provincias": [
                    {"codigo": code, "nombre": pname, "escanos": prov_seats[code],
                     "censo": metrics.get("Electores", {}).get(code),
                     "votantes": metrics.get("Votantes", {}).get(code),
                     "votos_validos": metrics.get("Votos válidos", {}).get(code),
                     "votos_blancos": metrics.get("Votos en blanco", {}).get(code),
                     "votos_nulos": metrics.get("Votos nulos", {}).get(code)}
                    for (ci, code, pname) in provs],
                "partidos": sorted(
                    [{"partido_id": pid, "votos": v["votos"], "escanos": v["escanos"]}
                     for pid, v in pt.items()],
                    key=lambda x: (-x["escanos"], -x["votos"])),
                "totales": {k: sum(metrics.get(kk, {}).get(c, 0) or 0 for _, c, _ in provs)
                            for k, kk in [("censo", "Electores"), ("votantes", "Votantes"),
                                          ("votos_validos", "Votos válidos"),
                                          ("votos_blancos", "Votos en blanco"),
                                          ("votos_nulos", "Votos nulos")]},
            })
            print(f"  {date}  escanos={val['escanos_totales']}  "
                  f"provincias OK={val['resumen']['provincias_ok']}/{val['resumen']['provincias_total']}  "
                  f"(a)={val['a_sum_escanos_ok']} (b)={val['b_dhondt_ok']} (c)={val['c_cargos_ok']}")

    partidos = {pid: {"nombre": sorted(m["nombres"], key=len, reverse=True)[0] if m["nombres"] else pid,
                      "siglas": sorted(m["siglas"]), "auto": m["auto"]}
                for pid, m in partidos_used.items()}
    with open(os.path.join(args.out, "partidos.json"), "w", encoding="utf-8") as f:
        json.dump({"version": "1.0", "partidos": partidos}, f, ensure_ascii=False, indent=1)

    with open(os.path.join(args.out, "elections.json"), "w", encoding="utf-8") as f:
        json.dump({
            "generado": _dt.datetime.now(_dt.timezone.utc).isoformat(),
            "fuente": {
                "ministerio": "https://descargas.interior.gob.es/datasets/resultados_electorales/",
                "portal_solicitado": "https://infoelectoral.interior.gob.es/ (bloqueado desde el host)",
                "ficheros": [CONGRESO_XLSX, CARGOS_XLSX],
            },
            "elecciones": elections,
        }, f, ensure_ascii=False, indent=1)

    with open(os.path.join(args.out, "validation.json"), "w", encoding="utf-8") as f:
        json.dump({"generado": _dt.datetime.now(_dt.timezone.utc).isoformat(),
                   "validaciones": validations}, f, ensure_ascii=False, indent=1)

    print(f"\nregistros provinciales: {n_rec}  ·  elecciones: {len(elections)}")
    print(f"salida en {args.out}/")


if __name__ == "__main__":
    main()
